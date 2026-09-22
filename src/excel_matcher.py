import argparse
import os
from pathlib import Path
import re
import sys
from bs4 import BeautifulSoup
from dotenv import load_dotenv
import openpyxl
import xlrd

# Ensure UTF-8 output encoding for terminal on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from src.telegram_bot import send_telegram_message
from src.gmail_client import get_gmail_service, list_recent_messages, get_message_details, download_attachment

TEMP_ATTACHMENTS_DIR = BASE_DIR / "data" / "temp_attachments"


def get_mailbox2_identifiers():
    """Load registration number and NeoPat ID from .env."""
    load_dotenv(dotenv_path=BASE_DIR / ".env", override=True)
    reg_no = os.getenv("REGISTRATION_NUMBER", "").strip()
    neopat_id = os.getenv("NEOPAT_ID", "").strip()
    return reg_no, neopat_id


def search_identifiers_in_text(text, reg_no, neopat_id):
    """Search for identifiers in subject or body text case-insensitively."""
    matched = set()
    if not text:
        return matched

    lower_text = text.lower()
    if reg_no and reg_no.lower() in lower_text:
        matched.add("reg_no")
    if neopat_id and neopat_id.lower() in lower_text:
        matched.add("neopat_id")

    return matched


def search_xlsx_file(file_path, reg_no, neopat_id):
    """Search every worksheet and every cell in an .xlsx file using openpyxl."""
    matched = set()
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        cell_str = str(cell).lower()
                        if reg_no and reg_no.lower() in cell_str:
                            matched.add("reg_no")
                        if neopat_id and neopat_id.lower() in cell_str:
                            matched.add("neopat_id")
        wb.close()
    except Exception as e:
        print(f"[Warning] Could not parse .xlsx file '{Path(file_path).name}': {type(e).__name__} - {e}")
    return matched


def search_xls_file(file_path, reg_no, neopat_id):
    """Search every worksheet and every cell in an .xls file using xlrd."""
    matched = set()
    try:
        wb = xlrd.open_workbook(file_path)
        for sheet_idx in range(wb.nsheets):
            sheet = wb.sheet_by_index(sheet_idx)
            for row_idx in range(sheet.nrows):
                for col_idx in range(sheet.ncols):
                    cell_val = sheet.cell_value(row_idx, col_idx)
                    if cell_val is not None:
                        cell_str = str(cell_val).lower()
                        if reg_no and reg_no.lower() in cell_str:
                            matched.add("reg_no")
                        if neopat_id and neopat_id.lower() in cell_str:
                            matched.add("neopat_id")
    except Exception as e:
        print(f"[Warning] Could not parse .xls file '{Path(file_path).name}': {type(e).__name__} - {e}")
    return matched


def process_excel_attachment(service, message_id, attachment_meta, reg_no, neopat_id):
    """
    Download an Excel attachment temporarily, parse with openpyxl/xlrd, and clean up.
    Returns a tuple of (matched_set, filename).
    """
    filename = attachment_meta.get("filename", "")
    attachment_id = attachment_meta.get("attachment_id")
    lower_fn = filename.lower()

    if not (lower_fn.endswith(".xlsx") or lower_fn.endswith(".xls")):
        return set(), filename

    TEMP_ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = re.sub(r'[\\/*?:"<>|]', "_", filename)
    temp_file = TEMP_ATTACHMENTS_DIR / f"{message_id}_{safe_name}"

    matched = set()
    try:
        file_bytes = download_attachment(service, message_id, attachment_id)
        with open(temp_file, "wb") as f:
            f.write(file_bytes)

        if lower_fn.endswith(".xlsx"):
            matched = search_xlsx_file(temp_file, reg_no, neopat_id)
        elif lower_fn.endswith(".xls"):
            matched = search_xls_file(temp_file, reg_no, neopat_id)
    except Exception as e:
        print(f"[Warning] Could not process attachment '{filename}': {type(e).__name__} - {e}")
    finally:
        if temp_file.exists():
            try:
                temp_file.unlink()
            except Exception:
                pass

    return matched, filename


def format_matched_identifiers_label(matched_set, reg_no, neopat_id):
    """
    Format human-readable label according to which identifiers matched:
      - Registration number
      - NeoPat ID
      - Both identifiers
    """
    has_reg = "reg_no" in matched_set
    has_neo = "neopat_id" in matched_set

    if has_reg and has_neo:
        return f"Both identifiers ({reg_no}, {neopat_id})"
    elif has_reg:
        return f"Registration number ({reg_no})"
    elif has_neo:
        return f"NeoPat ID ({neopat_id})"
    return None


def extract_company_name_mailbox2(subject, body="", attachment_names=None):
    """
    Extract company name from subject, body, or attachment filenames.
    Falls back to 'Company name could not be determined'.
    """
    subject = subject or ""

    # 1. Primary subject patterns
    subject_patterns = [
        r"^([A-Za-z0-9\s&.,'()\-]+?)\s*[-–—:]\s*(?:online\s+test|next\s+round|shortlist|placement|interview|assessment|selection)",
        r"^([A-Za-z0-9\s&.,'()\-]+?)\s+(?:next\s+round|next\s+level|selection\s+process|shortlist|online\s+test|technical\s+interview|dream\s+core|placement\s+drive)",
        r"eligible\s+for\s+(?:this\s+company\s*[:\-]\s*)?([A-Za-z0-9\s&.,'()\-]+?)(?:\s+(?:placement|drive|recruitment|process|batch)|[.!\n\r]|$)",
        r"\bcompany\s*[:\-]\s*([A-Za-z0-9\s&.,'()\-]+?)(?:\s+(?:placement|drive|recruitment|process|batch)|[.!\n\r]|$)"
    ]

    for pat in subject_patterns:
        m = re.search(pat, subject, re.IGNORECASE)
        if m:
            comp = m.group(1).strip().strip(":.,-")
            comp_lower = comp.lower()
            if (
                comp
                and len(comp) > 1
                and comp_lower not in ("kind attention", "urgent", "important", "fwd", "re", "security alert", "update", "notice", "announcement")
                and not comp_lower.startswith(("shortlist", "list of", "candidates", "students", "eligible"))
            ):
                return comp

    # 2. Check attachment filenames (e.g. "Exxonmobil interview shortlist.xlsx" -> "Exxonmobil")
    if attachment_names:
        for fname in attachment_names:
            att_m = re.search(r"^([A-Za-z0-9\s&.,'()\-]+?)\s+(?:additional\s+shortlist|online\s+assessment\s+shortlist|interview\s+shortlist|shortlist)", fname, re.IGNORECASE)
            if att_m:
                comp = att_m.group(1).strip()
                if comp and len(comp) > 1:
                    return comp

    # 3. Check body text
    body_patterns = [
        r"(?:placement\s+drive\s+with|drive\s+with|interview\s+with|shortlisted\s+for|company\s*[:\-])\s*([A-Za-z0-9\s&.,'()\-]+?)(?:\s+(?:placement|drive|recruitment|process|batch)|[.!\n\r]|$)"
    ]
    if body:
        for pat in body_patterns:
            m = re.search(pat, body, re.IGNORECASE)
            if m:
                comp = m.group(1).strip().strip(":.,-")
                comp_lower = comp.lower().strip()
                if (
                    comp
                    and len(comp) > 1
                    and comp_lower not in (
                        "the", "a", "an", "this", "this company", "the technical interview",
                        "technical interview", "interview", "the interview", "the online test",
                        "the test", "the selection process", "selection process", "next round",
                        "the next round", "the process"
                    )
                    and not comp_lower.startswith(("the interview", "technical interview", "interview", "shortlist", "selection", "the technical"))
                ):
                    return comp

    return "Company name could not be determined"


def extract_context(text, subject=""):
    """
    Extract useful context fields from email subject and body:
    - Application deadline
    - Test or interview date
    - Venue
    - Reporting time
    - Next steps
    - Important links
    - Contact details
    """
    combined = f"{subject}\n{text}"
    context = {}

    # 1. Venue
    venue_m = re.search(r'(?:[@]|venue\s*[:\-]|at\s+venue\s*[:\-])\s*([A-Za-z0-9\s,–\-\(\)\./]{3,80}?)(?:\*|\n|\r|$)', combined, re.IGNORECASE)
    if venue_m:
        venue = venue_m.group(1).strip().strip("*,.")
        if venue and len(venue) > 2:
            context["Venue"] = venue

    # 2. Test / Interview Date
    date_m = re.search(r'(?:scheduled\s+(?:from|on)|date\s*[:\-])\s*([0-9]{2}[\-\/][0-9]{2}[\-\/][0-9]{4}(?:\s+to\s+[0-9]{2}[\-\/][0-9]{2}[\-\/][0-9]{4})?|\([0-9]{2}[\-\/][0-9]{2}[\-\/][0-9]{4}\))', combined, re.IGNORECASE)
    if date_m:
        context["Test/Interview Date"] = date_m.group(1).strip("()")

    # 3. Reporting Time
    time_m = re.search(r'(?:by|at|reporting\s+time\s*[:\-])\s*([0-9]{1,2}(?::[0-9]{2})?\s*(?:am|pm))', combined, re.IGNORECASE)
    if time_m:
        context["Reporting Time"] = time_m.group(1).strip()

    # 4. Application Deadline
    deadline_m = re.search(r'(?:deadline|last\s+date|closes\s+on|registration\s+closes)\s*[:\-]?\s*([A-Za-z0-9\s,–\-\/:\(\)]{5,50}?)(?:\.|\n|\r|$)', combined, re.IGNORECASE)
    if deadline_m:
        deadline = deadline_m.group(1).strip().strip("*,.")
        if deadline:
            context["Application Deadline"] = deadline

    # 5. Next steps / Instructions
    steps = []
    lower_comb = combined.lower()
    if "technical interview" in lower_comb:
        steps.append("Technical Interview")
    elif "interview" in lower_comb:
        steps.append("Interview scheduled")
    if "online test" in lower_comb or "online assessment" in lower_comb:
        steps.append("Online assessment / test")
    if "resume" in lower_comb:
        steps.append("Carry printed copies of resume")
    if "formal" in lower_comb:
        steps.append("Formal attire mandatory")
    if "laptop" in lower_comb:
        steps.append("Bring personal laptop")
    if steps:
        context["Next Steps"] = "; ".join(steps)

    # 6. Important Links
    links = re.findall(r'https?://[^\s<>\"\'\)]+', combined)
    if links:
        seen_links = []
        for l in links:
            if l not in seen_links and "google.com" not in l:
                seen_links.append(l)
        if seen_links:
            context["Important Links"] = ", ".join(seen_links[:2])

    # 7. Contact Details
    emails = re.findall(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', combined)
    if emails:
        seen_emails = []
        for e in emails:
            if e.lower() not in seen_emails and "students.cdc" not in e.lower() and "vitapstudent" not in e.lower():
                seen_emails.append(e)
        if seen_emails:
            context["Contact Details"] = ", ".join(seen_emails[:2])

    return context


# Strong non-placement patterns (subject or sender) to filter out security alerts, logins, bills, etc.
NON_PLACEMENT_SUBJECT_PATTERNS = [
    r"\bsecurity\s+alert\b",
    r"\bcritical\s+security\s+alert\b",
    r"\bnew\s+sign-in\b",
    r"\bsign-in\s+on\s+your\s+google\s+account\b",
    r"\bgoogle\s+account\b",
    r"\bpassword\s+(?:reset|changed?)\b",
    r"\bverification\s+code\b",
    r"\bverify\s+your\s+account\b",
    r"\btwo-step\s+verification\b",
    r"\b2-step\s+verification\b",
    r"\baccount\s+recovery\b",
    r"\bbilling\s+notification\b",
    r"\bpayment\s+(?:for\s+.*?\s+)?is\s+successful\b",
    r"\bpayment\s+successful\b",
    r"\bquarterly\s+reminder\b",
    r"\bterms\s+of\s+service\b",
    r"\bprivacy\s+policy\b",
    r"\bmanaging\s+your\s+data\b",
]

# Placement signal keywords
PLACEMENT_KEYWORDS = [
    "placement", "campus placement", "placement drive", "placements",
    "recruitment", "recruiter", "recruitment drive", "campus recruitment", "recruitment process",
    "shortlist", "shortlisted", "shortlisting",
    "eligible", "eligibility",
    "technical interview", "hr interview", "interview shortlist", "interview schedule", "interview",
    "selection process", "selection round", "selection status", "selection",
    "next round", "next level",
    "assessment", "coding assessment", "online assessment", "aptitude test", "online test",
    "test/interview", "test / interview",
    "company registration", "company drive",
    "cdc", "career development centre", "career development center", "placement cell",
    "pre-placement", "job description"
]

PLACEMENT_SENDER_PATTERNS = [
    r"\bcdc(?:info)?\b",
    r"students\.cdc",
    r"placement",
    r"career",
    r"recruitment"
]


def classify_placement_relevance(
    subject,
    body="",
    sender="",
    attachment_names=None,
    matched_in_attachment=False
):
    """
    Determine whether an email that matched student identifiers is genuinely placement-related.
    Returns a tuple of (is_placement: bool, reason: str).

    Distinguishes legitimate placement opportunities and shortlists from unrelated emails
    (such as Google security alerts, sign-in alerts, billing receipts, or general university notices).
    """
    subject_str = (subject or "").strip()
    body_str = (body or "").strip()
    sender_str = (sender or "").strip()
    attachment_names = attachment_names or []

    subject_lower = subject_str.lower()
    body_lower = body_str.lower()
    sender_lower = sender_str.lower()
    att_str_lower = " ".join(attachment_names).lower()
    combined_text = f"{subject_lower} {body_lower} {att_str_lower}"

    # 1. Check for disqualifying non-placement subject patterns
    for pat in NON_PLACEMENT_SUBJECT_PATTERNS:
        if re.search(pat, subject_lower):
            return False, f"Disqualified: Matched non-placement alert pattern in subject ('{pat}')"

    # 2. Check for general administrative university notifications without placement keywords
    if re.search(r"^(?:university|campus|general|hostel|library|transport)\s+notification\b", subject_lower):
        has_placement = any(kw in combined_text for kw in PLACEMENT_KEYWORDS)
        if not has_placement:
            return False, "Disqualified: General notification without recruitment or placement context"

    # 3. Collect positive placement signals
    signals = []

    # Check subject (high-confidence signal)
    for kw in PLACEMENT_KEYWORDS:
        if kw in subject_lower:
            signals.append(f"subject contains '{kw}'")

    # Check attachment names or matched Excel attachment
    if matched_in_attachment:
        signals.append("identifier matched inside Excel attachment")
    for kw in ["shortlist", "interview", "eligible", "selection", "placement", "drive", "assessment"]:
        if kw in att_str_lower:
            signals.append(f"attachment name contains '{kw}'")

    # Check sender authority (CDC / Placement Office)
    for pat in PLACEMENT_SENDER_PATTERNS:
        if re.search(pat, sender_lower):
            signals.append(f"sender matched placement authority ('{pat}')")
            break

    # Check body for placement terms
    for kw in PLACEMENT_KEYWORDS:
        if kw in body_lower:
            if len(signals) < 5:
                signals.append(f"body contains '{kw}'")

    # 4. Decision logic
    # An email is placement-related if positive signals are found
    if signals:
        return True, f"Placement signals detected: {', '.join(signals[:3])}"

    return False, "No credible placement or recruitment context found"


def format_mailbox2_alert(
    company_name,
    source_mailbox,
    sender,
    subject,
    date,
    matched_identifier,
    attachment_name=None,
    context_dict=None
):
    """
    Format Telegram notification containing:
    - Alert heading
    - Company name
    - Source mailbox
    - Sender
    - Subject
    - Date
    - Matched identifier
    - Attachment name, if applicable
    - Useful context
    """
    lines = [
        "🎯 Placement Shortlist Alert! 🎯",
        "",
        f"🏢 Company: {company_name}",
        f"📬 Source Mailbox: {source_mailbox}",
        f"👤 Sender: {sender}",
        f"📌 Subject: {subject}",
        f"📅 Date: {date}",
        f"🆔 Matched Identifier: {matched_identifier}",
    ]

    if attachment_name:
        lines.append(f"📎 Attachment: {attachment_name}")

    if context_dict:
        lines.append("")
        lines.append("📋 Useful Context:")
        for k, v in context_dict.items():
            lines.append(f"• {k}: {v}")

    return "\n".join(lines)


def process_mailbox2_email(
    service,
    msg_details,
    source_mailbox=None,
    send_alert=True
):
    """
    Inspect an email from Mailbox 2 for Registration Number and NeoPat ID.
    Searches:
      1. Subject
      2. Body
      3. All .xlsx and .xls attachments across all sheets and cells
    Sends Telegram alert if matched.
    Returns a result dict.
    """
    source_mailbox = source_mailbox or os.getenv("MAILBOX2", "Mailbox 2")
    reg_no, neopat_id = get_mailbox2_identifiers()

    subject = msg_details.get("subject", "")
    body = msg_details.get("body", "")
    sender = msg_details.get("sender", "(Unknown Sender)")
    date = msg_details.get("date", "(Unknown Date)")
    attachments = msg_details.get("attachments", [])
    msg_id = msg_details.get("id")

    # Clean text from body
    clean_body = ""
    if body:
        try:
            clean_body = BeautifulSoup(body, "html.parser").get_text(separator=" ")
        except Exception:
            clean_body = body

    all_matches = set()
    matched_attachment_name = None

    # 1. Search in subject
    subject_matches = search_identifiers_in_text(subject, reg_no, neopat_id)
    all_matches.update(subject_matches)

    # 2. Search in body
    body_matches = search_identifiers_in_text(clean_body, reg_no, neopat_id)
    all_matches.update(body_matches)

    # 3. Search in Excel attachments
    att_names = [a.get("filename", "") for a in attachments]
    for att in attachments:
        fn = att.get("filename", "")
        if fn.lower().endswith(".xlsx") or fn.lower().endswith(".xls"):
            att_matches, parsed_fn = process_excel_attachment(
                service,
                msg_id,
                att,
                reg_no,
                neopat_id
            )
            if att_matches:
                all_matches.update(att_matches)
                matched_attachment_name = parsed_fn

    if not all_matches:
        return {
            "matched": False,
            "identifier_matched": False,
            "is_placement": False,
            "classification": "NON-MATCH",
            "company": None,
            "matched_identifier": None,
            "alert_sent": False
        }

    # Identifier matched!
    matched_label = format_matched_identifiers_label(all_matches, reg_no, neopat_id)
    matched_in_attachment = matched_attachment_name is not None

    # Step 2: Classify placement relevance
    is_placement, classification_reason = classify_placement_relevance(
        subject=subject,
        body=clean_body,
        sender=sender,
        attachment_names=att_names,
        matched_in_attachment=matched_in_attachment
    )

    if not is_placement:
        # Non-placement email that merely contained the identifier
        return {
            "matched": False,
            "identifier_matched": True,
            "is_placement": False,
            "classification": "NON-PLACEMENT",
            "reason": classification_reason,
            "company": None,
            "matched_identifier": matched_label,
            "attachment": matched_attachment_name,
            "alert_sent": False,
            "alert_message": None
        }

    # Legitimate placement email! Extract company and context
    company_name = extract_company_name_mailbox2(subject, clean_body, att_names)
    context_dict = extract_context(clean_body, subject=subject)

    alert_message = format_mailbox2_alert(
        company_name=company_name,
        source_mailbox=source_mailbox,
        sender=sender,
        subject=subject,
        date=date,
        matched_identifier=matched_label,
        attachment_name=matched_attachment_name,
        context_dict=context_dict
    )

    if send_alert:
        print(f"[Mailbox 2] 🎯 Match found! Company: '{company_name}', Identifier: {matched_label}. Sending Telegram alert...")
        success = send_telegram_message(alert_message)
    else:
        print(f"[Mailbox 2] 🎯 Match found! Company: '{company_name}', Identifier: {matched_label}. (Dry-run mode: Alert not sent)")
        success = False

    return {
        "matched": True,
        "identifier_matched": True,
        "is_placement": True,
        "classification": "PLACEMENT",
        "reason": classification_reason,
        "company": company_name,
        "matched_identifier": matched_label,
        "attachment": matched_attachment_name,
        "alert_sent": success,
        "alert_message": alert_message
    }


def diagnostic_scan_mailbox2(max_emails=5, send_alert=False):
    """
    Safely scan recent Mailbox 2 emails and Excel attachments without mutating polling state.
    """
    reg_no, neopat_id = get_mailbox2_identifiers()
    source_mb = os.getenv("MAILBOX2", "Mailbox 2")
    print("=" * 65)
    print(f"Diagnostic Scan: Mailbox 2 ({source_mb})")
    print(f"Target Registration Number: {reg_no or '[Not configured]'}")
    print(f"Target NeoPat ID:           {neopat_id or '[Not configured]'}")
    print(f"Live Telegram Alerts:       {send_alert}")
    print("=" * 65)

    service = get_gmail_service(mailbox_id=2, interactive=False)
    if not service:
        print("Mailbox 2 service could not be initialized.")
        return

    messages = list_recent_messages(service, max_results=max_emails)
    print(f"Inspecting {len(messages)} recent message(s)...\n")

    matched_count = 0
    for idx, msg_meta in enumerate(messages, start=1):
        details = get_message_details(service, msg_meta["id"])
        att_names = [a["filename"] for a in details.get("attachments", [])]
        print(f"[{idx}] Subject: {details['subject']}")
        print(f"    From: {details['sender']}")
        if att_names:
            print(f"    Attachments: {att_names}")

        result = process_mailbox2_email(
            service,
            details,
            send_alert=send_alert
        )

        if result["matched"]:
            matched_count += 1
            print(f"    👉 PLACEMENT MATCH DETECTED! Company: '{result['company']}', Match: {result['matched_identifier']}")
            if result.get("attachment"):
                print(f"       Matched in attachment: {result['attachment']}")
        elif result.get("identifier_matched"):
            print(f"    ℹ️ Identifier matched ({result['matched_identifier']}), but classified as NON-PLACEMENT ({result.get('reason')}).")
        else:
            print(f"    ❌ Neither identifier ({reg_no} / {neopat_id}) found.")
        print()

    print(f"Diagnostic scan complete. Matched {matched_count} out of {len(messages)} email(s).")


def test_mailbox2_simulation():
    """Simulate a detection event with an Excel shortlist and deliver a live test notification to Telegram."""
    sample_email = {
        "id": "test-sample-mailbox2",
        "sender": "'No Reply CDC Info' via 2027 CDC <students.cdc2027@vitap.ac.in>",
        "subject": "Exxonmobil next round of selection process (Technical Interview) is scheduled on (18-09-2026) 08:00 AM @ CDC 717, SJT 7th floor - VIT Vellore",
        "date": "Thu, 17 Sep 2026 18:41:49 +0530",
        "body": """*Exxonmobil next round of selection process (Technical Interview) is scheduled on (18-09-2026) 08:00 AM @ CDC 717, SJT 7th floor - VIT Vellore*
*Please find the attached list of shortlisted students who cleared the group discussion*
*Students are informed to Carry 3 sets of their printed resume and college ID card.*
*Wearing Formals is mandatory.*
*Kindly attend the selection process without fail.*
*Additionally, please instruct all students to bring their personal laptops*
Visit https://vit.ac.in/placement for updates or contact cdc@vit.ac.in""",
        "attachments": [
            {"filename": "Exxonmobil interview shortlist.xlsx", "attachment_id": "mock_id"}
        ]
    }

    # Simulate matched identifier
    reg_no, neopat_id = get_mailbox2_identifiers()
    matched_set = {"reg_no", "neopat_id"}
    matched_label = format_matched_identifiers_label(matched_set, reg_no, neopat_id)
    company_name = extract_company_name_mailbox2(sample_email["subject"], sample_email["body"], ["Exxonmobil interview shortlist.xlsx"])
    context_dict = extract_context(sample_email["body"], subject=sample_email["subject"])

    alert = format_mailbox2_alert(
        company_name=company_name,
        source_mailbox=os.getenv("MAILBOX2", "Mailbox 2"),
        sender=sample_email["sender"],
        subject=sample_email["subject"],
        date=sample_email["date"],
        matched_identifier=matched_label,
        attachment_name="Exxonmobil interview shortlist.xlsx",
        context_dict=context_dict
    )

    print("=== Testing Live Mailbox 2 Shortlist Alert to Telegram ===")
    print(alert)
    print("\nDispatching alert to Telegram...")
    success = send_telegram_message(alert)
    if success:
        print("\n✅ Verification successful! Telegram shortlist alert was delivered.")
    else:
        print("\n❌ Notification delivery failed.")
    return success


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="PlacementMonitor Mailbox 2 Excel Matcher")
    parser.add_argument("--scan", action="store_true", help="Run safe diagnostic scan on Mailbox 2")
    parser.add_argument("--send-alert", action="store_true", help="Send live Telegram alerts during --scan if match found")
    parser.add_argument("--test-alert", action="store_true", help="Simulate and deliver a live Mailbox 2 Telegram alert")
    args = parser.parse_args()

    if args.test_alert:
        test_mailbox2_simulation()
    elif args.scan:
        diagnostic_scan_mailbox2(send_alert=args.send_alert)
    else:
        print("PlacementMonitor Mailbox 2 Excel Matcher ready.")
        print("Usage:")
        print("  .\\.venv\\Scripts\\python.exe src/excel_matcher.py --scan")
        print("  .\\.venv\\Scripts\\python.exe src/excel_matcher.py --test-alert")
